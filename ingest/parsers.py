"""
Format-specific parsers.

Each parse_* function takes a file path and returns a DocumentJSON object.
parse_any() is the single entrypoint that dispatches based on file extension.

Supported: .pdf, .docx, .csv, .xlsx, .txt, .md
"""

import os
from typing import List

import pandas as pd
from pypdf import PdfReader
from docx import Document as DocxDocument

from .schema import DocumentJSON, ContentBlock


def parse_pdf(path: str) -> DocumentJSON:
    reader = PdfReader(path)
    blocks: List[ContentBlock] = []

    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        # Split into paragraphs on blank lines; fall back to whole page text.
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        if not paragraphs and text.strip():
            paragraphs = [text.strip()]
        for para in paragraphs:
            blocks.append(ContentBlock(type="paragraph", text=para, page=page_num))

    metadata = {
        "pages": len(reader.pages),
        "author": (reader.metadata.author if reader.metadata else None),
        "title": (reader.metadata.title if reader.metadata else None),
    }
    return DocumentJSON(
        source_file=os.path.basename(path),
        file_type="pdf",
        metadata=metadata,
        content=blocks,
    )


def parse_docx(path: str) -> DocumentJSON:
    doc = DocxDocument(path)
    blocks: List[ContentBlock] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "").lower()
        if "heading" in style_name:
            # style names look like "Heading 1", "Heading 2", ...
            level = "".join(ch for ch in style_name if ch.isdigit())
            blocks.append(
                ContentBlock(
                    type="heading",
                    text=text,
                    level=int(level) if level else 1,
                )
            )
        else:
            blocks.append(ContentBlock(type="paragraph", text=text))

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            row_text = " | ".join(cells)
            if row_text.strip(" |"):
                blocks.append(
                    ContentBlock(
                        type="table_row",
                        text=row_text,
                        row=r_idx,
                        extra={"table_index": t_idx},
                    )
                )

    return DocumentJSON(
        source_file=os.path.basename(path),
        file_type="docx",
        metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
        content=blocks,
    )


def _parse_tabular(path: str, file_type: str) -> DocumentJSON:
    if file_type == "csv":
        df = pd.read_csv(path)
    else:  # xlsx
        df = pd.read_excel(path)

    blocks: List[ContentBlock] = []
    columns = list(df.columns)

    for row_idx, row in df.iterrows():
        # Represent each row as "col: value | col: value ..." so it reads
        # naturally when retrieved and shown to the LLM.
        row_text = " | ".join(f"{col}: {row[col]}" for col in columns)
        blocks.append(ContentBlock(type="table_row", text=row_text, row=int(row_idx)))

    return DocumentJSON(
        source_file=os.path.basename(path),
        file_type=file_type,
        metadata={"columns": columns, "num_rows": len(df)},
        content=blocks,
    )


def parse_csv(path: str) -> DocumentJSON:
    return _parse_tabular(path, "csv")


def parse_xlsx(path: str) -> DocumentJSON:
    return _parse_tabular(path, "xlsx")


def parse_text(path: str) -> DocumentJSON:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        raw = f.read()

    blocks = [
        ContentBlock(type="paragraph", text=p.strip())
        for p in raw.split("\n\n")
        if p.strip()
    ]
    ext = os.path.splitext(path)[1].lstrip(".").lower() or "txt"
    return DocumentJSON(
        source_file=os.path.basename(path),
        file_type=ext,
        metadata={"chars": len(raw)},
        content=blocks,
    )


_DISPATCH = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".csv": parse_csv,
    ".xlsx": parse_xlsx,
    ".xls": parse_xlsx,
    ".txt": parse_text,
    ".md": parse_text,
}


def parse_any(path: str) -> DocumentJSON:
    ext = os.path.splitext(path)[1].lower()
    if ext not in _DISPATCH:
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported: {list(_DISPATCH.keys())}"
        )
    return _DISPATCH[ext](path)
